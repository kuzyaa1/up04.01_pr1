from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out = Path(r'C:\Users\Viki\Desktop\ТЕХ\4 курс\УП\Отчет_Практическая_работа_1.docx')
doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2)
sec.bottom_margin = Cm(2)
sec.left_margin = Cm(3)
sec.right_margin = Cm(1.5)
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
style.font.size = Pt(14)
style.paragraph_format.first_line_indent = Cm(1.25)
style.paragraph_format.line_spacing = 1.5

def text(t='', bold=False, size=None, align=None, indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    if not indent: p.paragraph_format.first_line_indent = Cm(0)
    if align is not None: p.alignment = align
    r=p.add_run(t); r.bold=bold
    r.font.name='Times New Roman'; r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    if size: r.font.size=Pt(size)
    return p

def heading(t, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    r=p.add_run(t); r.bold=True; r.font.name='Times New Roman'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); r.font.size=Pt(14)
    return p

def cell_text(cell, value, bold=False):
    cell.text = ''
    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
    r=p.add_run(value); r.bold=bold; r.font.name='Times New Roman'; r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman'); r.font.size=Pt(11)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        cell_text(t.rows[0].cells[i], h, True); shade(t.rows[0].cells[i], 'D9E2F3')
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row): cell_text(cells[i], val)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

# Title page
for line in [
    'Министерство науки и высшего образования Российской Федерации',
    'Образовательная организация',
    '', '',
    'ОТЧЁТ',
    'по практической работе № 1',
    '«Маршрутизация, формы и URL как состояние»',
    '', '', '',
    'Выполнил(а): _______________________________',
    'Группа: ____________________________________',
    'Проверил(а): _______________________________',
    '', '',
    '2026'
]:
    p=text(line, bold=line in ('ОТЧЁТ','по практической работе № 1','«Маршрутизация, формы и URL как состояние»'), align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    if line.startswith('Выполнил') or line.startswith('Группа') or line.startswith('Проверил'):
        p.alignment=WD_ALIGN_PARAGRAPH.RIGHT

doc.add_page_break()
heading('1. Цель работы')
text('Разработать web-приложение на Flutter из нескольких экранов, освоить маршрутизацию через адресную строку, передачу параметров через URL, обработку прямого перехода по ссылке и проверку пользовательского ввода с помощью Form.')
heading('2. Использованные средства разработки')
text('Для выполнения работы использованы Flutter 3.27.2, Dart 3.6.1, Visual Studio Code с расширениями Dart и Flutter, браузер Google Chrome. Для маршрутизации подключён пакет go_router, для управления темой — provider, для сохранения настроек — shared_preferences.')
heading('3. Описание разработанного приложения')
text('Разработано web-приложение «Калькулятор и конвертер». Главная страница содержит текстовые кнопки перехода к калькулятору и конвертеру валют. Все графические значки в интерфейсе удалены; элементы управления содержат только текст.')
text('Калькулятор принимает два числа и выбранную операцию: сложение, вычитание, умножение или деление. После проверки формы выполняется переход на отдельный экран результата. Данные не передаются через конструктор виджета или extra: значения записываются в query-параметры адреса.')
text('Конвертер валют содержит два выпадающих списка и поле суммы. Поддерживаются валюты RUB, USD, EUR, CNY, KZT и BYN. Курсы заданы константами в файле логики. Выбранная валютная пара сохраняется в localStorage через shared_preferences и восстанавливается при новом открытии экрана.')
heading('4. Структура проекта')
table(['Каталог или файл', 'Назначение'], [
    ('lib/main.dart', 'Точка входа, MaterialApp.router, тема и usePathUrlStrategy().'),
    ('lib/router.dart', 'Конфигурация маршрутов go_router и экран 404.'),
    ('lib/screens/', 'Экраны главной страницы, калькулятора, конвертера, результата и ошибки.'),
    ('lib/logic/calculator.dart', 'Чистая логика арифметических вычислений.'),
    ('lib/logic/currency.dart', 'Курсы валют и логика конвертации.'),
    ('lib/core/format.dart', 'Округление и форматирование результата.'),
    ('lib/state/theme_notifier.dart', 'Состояние светлой/тёмной темы и его сохранение.'),
    ('test/', 'Модульные тесты калькулятора, конвертера и форматирования.')
], [5.4, 10.1])
heading('5. Схема маршрутов')
table(['Адрес', 'Экран и назначение'], [
    ('/', 'Главная страница с переходами в разделы.'),
    ('/calculator', 'Форма калькулятора.'),
    ('/calculator/result?a=...&op=...&b=...', 'Экран результата вычисления. Все данные извлекаются из URL.'),
    ('/converter', 'Форма конвертера валют.'),
    ('/converter/result?from=...&to=...&amount=...', 'Экран результата конвертации. Все данные извлекаются из URL.'),
    ('Любой неизвестный адрес, например /qwerty', 'Экран 404, создаваемый errorBuilder.')
], [7.2, 8.3])
text('В точке входа вызывается usePathUrlStrategy(), поэтому в адресах отсутствует символ решётки. Навигация между разделами выполняется методом context.go(), который синхронизирует адрес приложения с историей браузера.')
heading('6. Передача данных через URL')
text('После успешной проверки формы калькулятора значения кодируются и записываются в строку запроса. Например, вычисление 10 + 5 открывает адрес /calculator/result?a=10&op=%2B&b=5. Символ «+» обязательно кодируется функцией Uri.encodeComponent(), так как в строке запроса он иначе интерпретируется как пробел.')
text('Экран результата читает параметры через state.uri.queryParameters и повторно выполняет вычисление. Поэтому прямое открытие ссылки в новой вкладке, обновление страницы и ручной ввод адреса обрабатываются корректно. При отсутствии параметра, нечисловом значении, неизвестной операции или делении на ноль выводится текстовая ошибка и кнопка возврата к форме.')
heading('7. Формы и валидация')
text('Обе формы реализованы с помощью Form, GlobalKey<FormState> и TextFormField. Используется режим AutovalidateMode.onUserInteraction. Для числовых полей проверяются пустой ввод и корректность преобразования в double; для суммы дополнительно запрещено отрицательное значение. Выпадающие списки валют также участвуют в проверке формы.')
heading('8. Дополнительные возможности')
table(['Требование', 'Реализация'], [
    ('Адаптивная вёрстка', 'Формы ограничены ConstrainedBox с максимальной шириной 420 px и помещены в SingleChildScrollView. Интерфейс сохраняет работоспособность при ширине от 360 до 1920 px.'),
    ('Светлая и тёмная тема', 'ThemeNotifier хранит ThemeMode; выбор сохраняется в shared_preferences и восстанавливается после перезагрузки.'),
    ('Запоминание валют', 'Последняя пара валют сохраняется в localStorage через shared_preferences.'),
    ('Округление', 'Результат форматируется до разумного количества десятичных знаков, хвостовые нули удаляются.'),
    ('Тестирование', 'Логика вынесена из виджетов; написано 19 модульных проверок.')
], [5.3, 10.2])
heading('9. Проверка работы программы')
text('Команда flutter analyze выполнена без замечаний: No issues found. Команда flutter test успешно завершена: All tests passed, выполнено 19 тестов. Команда flutter build web успешно сформировала web-сборку в каталоге build/web.')
text('Тесты проверяют арифметические операции, деление на ноль, нечисловой и отсутствующий ввод, неизвестную операцию, обработку десятичной запятой, прямую и обратную конвертацию, кросс-курс валют, недопустимую валюту, отрицательную сумму и округление результата.')
heading('10. Заключение')
text('В ходе работы создано Flutter Web-приложение с маршрутизацией на go_router, формами и валидацией. Состояние результата представлено в URL, поэтому ссылки можно открывать напрямую и обновлять без потери данных. Реализованы экран 404, обработка ошибочных параметров, адаптивная раскладка, сохранение темы и пары валют, а вычислительная логика покрыта модульными тестами.')

# Footer page numbering
for section in doc.sections:
    footer=section.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)

doc.save(out)
print(out)
